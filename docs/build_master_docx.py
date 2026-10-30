# -*- coding: utf-8 -*-
"""
Master Document Generator for NagarYojana (नगरयोजना)
Generates the comprehensive, production-grade master architectural blueprint,
feature catalog, REST API specification, database schema, security runbook,
and visual interface gallery for docs/overalldocx.docx.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = Document()

    # Configure 0.75-inch standard margins
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)

    # Palette Constants
    COLOR_TITLE_RED = RGBColor(0xDC, 0x14, 0x3C)     # Crimson Red
    COLOR_TITLE_GREEN = RGBColor(0x15, 0x69, 0x47)   # Forest Civic Green
    COLOR_SLATE_900 = RGBColor(0x0F, 0x17, 0x2A)     # Heading 1 & Table Header
    COLOR_BLUE_900 = RGBColor(0x1E, 0x3A, 0x8A)      # Heading 2
    COLOR_BODY_INK = RGBColor(0x1D, 0x1D, 0x1F)      # Apple Body Ink
    COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)         # Subtitles & Captions
    COLOR_CODE_TEXT = RGBColor(0xF1, 0xF5, 0xF9)     # Code Block Light Text

    # Helper: Set table borders
    def set_table_borders(table, color="E2E8F0", sz="4"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    # Typography Helpers
    def add_header_banner():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)

        r1 = p.add_run("Nagar")
        r1.font.name = "Arial"
        r1.font.size = Pt(32)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_TITLE_RED

        r2 = p.add_run("योजना")
        r2.font.name = "Arial"
        r2.font.size = Pt(32)
        r2.font.bold = True
        r2.font.color.rgb = COLOR_TITLE_GREEN

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(3)
        r_sub = p_sub.add_run("Complete Architectural Blueprint, System Implementation Guide & Production Runbook")
        r_sub.font.name = "Arial"
        r_sub.font.size = Pt(13)
        r_sub.font.bold = True
        r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

        p_meta = doc.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_meta.paragraph_format.space_before = Pt(0)
        p_meta.paragraph_format.space_after = Pt(14)
        r_meta = p_meta.add_run("Target Audience: Civic Administrators, Municipal Engineers, Students & Full-Stack Engineers | Written in Plain English | Updated October 30, 2026")
        r_meta.font.name = "Arial"
        r_meta.font.size = Pt(9.5)
        r_meta.font.color.rgb = COLOR_MUTED

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(15.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SLATE_900

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_BLUE_900

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_SLATE_900

    def add_p(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_BODY_INK
        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = COLOR_BODY_INK

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25)
        
        r_sym = p.add_run("• ")
        r_sym.font.name = "Arial"
        r_sym.font.size = Pt(10)
        r_sym.font.bold = True
        r_sym.font.color.rgb = COLOR_BLUE_900

        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Arial"
            r_b.font.size = Pt(10)
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_BODY_INK

        r_t = p.add_run(text)
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = COLOR_BODY_INK

    def add_callout(text, title=None, border_color="0066CC", bg_color="F8FAFC"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        cell = tbl.cell(0, 0)
        cell.width = Inches(7.0)

        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        tcPr.append(shd)

        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="140" w:type="dxa"/>'
            f'  <w:bottom w:w="140" w:type="dxa"/>'
            f'  <w:left w:w="200" w:type="dxa"/>'
            f'  <w:right w:w="200" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15

        if title:
            rt = p.add_run(f"📌 {title}\n")
            rt.font.name = "Arial"
            rt.font.size = Pt(10)
            rt.font.bold = True
            rt.font.color.rgb = COLOR_SLATE_900

        rtxt = p.add_run(text)
        rtxt.font.name = "Arial"
        rtxt.font.size = Pt(9.5)
        rtxt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        # Space after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def add_code_block(code_str, title=None):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        cell = tbl.cell(0, 0)
        cell.width = Inches(7.0)

        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
        tcPr.append(shd)

        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="38BDF8"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top w:w="120" w:type="dxa"/>'
            f'  <w:bottom w:w="120" w:type="dxa"/>'
            f'  <w:left w:w="180" w:type="dxa"/>'
            f'  <w:right w:w="180" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05

        if title:
            rt = p.add_run(f"// {title}\n")
            rt.font.name = "Consolas"
            rt.font.size = Pt(8.5)
            rt.font.bold = True
            rt.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)

        rc = p.add_run(code_str)
        rc.font.name = "Consolas"
        rc.font.size = Pt(8.5)
        rc.font.color.rgb = COLOR_CODE_TEXT

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def add_table_data(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, color="CBD5E1", sz="4")

        # Header formatting
        hdr_row = table.rows[0]
        for col_idx, header_text in enumerate(headers):
            cell = hdr_row.cells[col_idx]
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]

            tcPr = cell._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
            tcPr.append(shd)

            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                f'  <w:top w:w="100" w:type="dxa"/>'
                f'  <w:bottom w:w="100" w:type="dxa"/>'
                f'  <w:left w:w="120" w:type="dxa"/>'
                f'  <w:right w:w="120" w:type="dxa"/>'
                f'</w:tcMar>'
            )
            tcPr.append(tcMar)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(header_text)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Data rows formatting
        for r_idx, row_values in enumerate(rows_data):
            row = table.rows[r_idx + 1]
            bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"

            for col_idx, val in enumerate(row_values):
                cell = row.cells[col_idx]
                if col_widths and col_idx < len(col_widths):
                    cell.width = col_widths[col_idx]

                tcPr = cell._tc.get_or_add_tcPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
                tcPr.append(shd)

                tcMar = parse_xml(
                    f'<w:tcMar {nsdecls("w")}>'
                    f'  <w:top w:w="80" w:type="dxa"/>'
                    f'  <w:bottom w:w="80" w:type="dxa"/>'
                    f'  <w:left w:w="100" w:type="dxa"/>'
                    f'  <w:right w:w="100" w:type="dxa"/>'
                    f'</w:tcMar>'
                )
                tcPr.append(tcMar)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                run = p.add_run(str(val))
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_BODY_INK

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    def add_image_figure(image_filename, caption_text):
        img_path = os.path.join("temp_images", image_filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(6.6))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.name = "Arial"
            r_cap.font.size = Pt(9.5)
            r_cap.font.bold = True
            r_cap.font.color.rgb = COLOR_MUTED

    # ==========================================
    # BUILD DOCUMENT CONTENT
    # ==========================================

    add_header_banner()

    add_callout(
        "If you are evaluating the civic governance infrastructure of Biratnagar Metropolitan City or learning full-stack web engineering, this document provides an exhaustive blueprint of the Nagarयोजना platform. It explains the high-level layers, the 7-step automated credential email pipeline, the 3-attempt brute-force account lockout defense, the unified municipal login gateway with automatic 3-tier role redirection (Super Admin, Ward Executive, Contractor), the field milestone progress tracking system, the municipal maintenance break mode with instant cross-tab eviction, the zero-login citizen grievance reporting modal, the zero-storage in-browser dynamic bilingual engine, the complete REST API specification, the Firestore database schemas, and an actionable step-by-step production runbook.",
        title="Purpose of this Master Guide"
    )

    # ------------------------------------------
    # SECTION 1: EXECUTIVE OVERVIEW & CIVIC MISSION
    # ------------------------------------------
    add_h1("1. Executive Overview & Civic Mission")
    add_p(
        "is an open civic infrastructure monitoring and municipal governance platform designed for local governments in Nepal, currently configured and operating for Biratnagar Metropolitan City (विराटनगर महानगरपालिका) in Koshi Province.",
        bold_prefix="Nagarयोजना (नगरयोजना) "
    )
    add_p(
        "In traditional municipal operations across Nepal's 753 local bodies, a fundamental transparency and accountability gap exists between citizens and municipal authorities. Local residents frequently lack visibility into ward-level public works: they do not know what roads, drains, water supply systems, or community centers have been budgeted, who the appointed executing contractors are, what the planned completion deadlines are, or why a neighborhood project has stalled. Meanwhile, metropolitan administrators and ward committees struggle with manual paper records, unstructured citizen grievances, unverified field milestone reports, and slow credential provisioning."
    )
    add_p(
        "Nagarयोजना fundamentally solves these challenges by establishing a real-time, bi-directional digital bridge between citizens, municipal executives, and field construction contractors. Every municipal rupee and construction milestone is placed on a public screen, while administrative staff receive enterprise-grade workflows to govern, monitor, and deliver civic projects."
    )

    add_h2("1.1 Municipal Identity & Official Tagline")
    add_bullet("Official Tagline: 'Hamro Nagar, Aafno Sarkar: Ward's Tasks, On Your Screen.' (हाम्रो नगर, आफ्नो सरकार: वडाका काम, आफ्नै स्क्रिनमा)")
    add_bullet("Municipal Jurisdiction: Biratnagar Metropolitan City (विराटनगर महानगरपालिका), Koshi Province, Nepal — covering all 19 autonomous municipal wards (Ward 01 through Ward 19).")
    add_bullet("Production Live Website: https://nagaryojana.shrawankarki.com.np/")
    add_bullet("Production Live API Service: https://nagar-yojana.onrender.com")
    add_bullet("Principal Developer & Architect: Shrawan Karki (Portfolio: https://shrawankarki.com.np)")

    add_h2("1.2 Stakeholder Ecosystem")
    add_p("The platform serves four interconnected civic user personas:")
    add_bullet("Explore public infrastructure works across all 19 wards or filter by neighborhood; view live completion percentages, allocated budgets, and contractor details; file direct grievances without needing an account; switch dynamically between English and authentic Nepali.", bold_prefix="1. General Citizens & Taxpayers: ")
    add_bullet("Directly oversee projects within their respective ward; establish new infrastructure projects; provision and monitor field civic contractors (*Upabhokta Samiti*); inspect and resolve citizen grievances with a structured lifecycle audit.", bold_prefix="2. Ward Executives & Local Oversight Officers: ")
    add_bullet("Log directly into a streamlined mobile-friendly field workspace; report physical completion percentages (0–100%); set milestone statuses (Active, Completed, Delayed); submit engineering remarks and site logs.", bold_prefix="3. Civic Contractors & On-Site Engineers: ")
    add_bullet("Metropolitan Mayoral office and Chief Administrative Officers with complete oversight over all 19 wards; visual jurisdiction matrix; executive provisioning and credential generation; account lockout recovery; system-wide maintenance mode.", bold_prefix="4. Super Administrators & Central Municipal Leadership: ")

    # ------------------------------------------
    # SECTION 2: HIGH-LEVEL FULL-STACK ARCHITECTURE
    # ------------------------------------------
    add_h1("2. High-Level Architecture & Technical Foundation")
    add_p("Nagarयोजना is engineered as a decoupled, multi-cloud web application designed for high availability, zero cold-start disruption, instant synchronization, and sub-100 millisecond global asset delivery.")

    add_h2("2.1 Architectural Component Matrix")
    add_table_data(
        headers=["Architectural Layer", "Technology Stack", "Core Responsibility & Capabilities"],
        rows_data=[
            [
                "Frontend Client SPA",
                "React 18 + Vite 5 + Tailwind CSS 3 + Lucide Icons",
                "Renders the public transparency portal, interactive 19-ward filter bar, Apple-inspired project cards, citizen grievance modal, dynamic translator, and the Unified Municipal Gateway with 3-tier role workspaces."
            ],
            [
                "Backend RESTful API",
                "Node.js + Express 4 + Helmet + Dynamic CORS",
                "Enforces Role-Based Access Control (RBAC), sequential institutional ID calculation (WE-X-Y / CON-X-Y), zero-ambiguity credential generation, atomic rollbacks, grievance status transitions, and keep-alive health endpoints."
            ],
            [
                "Cloud NoSQL Database",
                "Google Firebase Cloud Firestore",
                "Stores real-time document collections: projects, users, grievances, maintenance settings, and failed login lockout records with atomic transaction guarantees and sub-second onSnapshot listener pushes."
            ],
            [
                "Identity & Access Management",
                "Google Firebase Authentication + Custom Claims",
                "Manages cryptographic JWT identity tokens, client-side browserSessionPersistence enforcement, and fine-grained custom claims (SUPER_ADMIN, WARD_EXEC, CONTRACTOR) decoded upon authentication."
            ],
            [
                "Transactional Email Engine",
                "Official Resend SDK + DNS Custom Domain",
                "Transmits DNS-verified (SPF, DKIM, DMARC) municipal HTML credential cards from admin@shrawankarki.com.np featuring copy-safe temporary password containers and pre-linked authentication buttons."
            ],
            [
                "Defense & Rate-Limiting",
                "IP-Bucket Limiter + 3-Attempt Lockout + 15m DOM Watcher",
                "Guards against credential brute-forcing with 30 req/min IP limits, 24-hour lockout after 3 consecutive failures, 1-click admin unlocking, and automated 15-minute DOM inactivity session termination."
            ],
            [
                "Multi-Cloud Production Hosting",
                "Cloudflare Pages (Edge) + Render (Node Container)",
                "Splits static SPA distribution across Cloudflare's global edge network (HTTP/3, sub-100ms) and dynamic API execution on Render's web service, connected seamlessly via dynamic environment detection."
            ],
            [
                "Continuous Keep-Alive",
                "UptimeRobot External Health Pinger",
                "Monitors root GET / and GET /health endpoints at 5-minute intervals, completely neutralizing free-tier server cold starts and ensuring zero startup delay for municipal field workers."
            ]
        ],
        col_widths=[Inches(1.8), Inches(1.8), Inches(3.4)]
    )

    add_h2("2.2 Apple-Inspired Civic Design Philosophy")
    add_p(
        "Civic governance interfaces are notorious for being cluttered, outdated, and intimidating. Nagarयोजना takes an entirely different design philosophy inspired by high-end consumer technology (documented in DESIGN.md):"
    )
    add_bullet("Minimalist Chrome: Interface chrome recedes into the background so public works, budgets, and milestones take center stage. No unnecessary gradients or distracting drop shadows.", bold_prefix="Clean Aesthetic: ")
    add_bullet("Utilizes Apple's San Francisco system typography hierarchy with tight negative letter-spacing for large headlines and generous line heights for body descriptions.", bold_prefix="Typography First: ")
    add_bullet("A single signature Action Blue (#0066CC) guides citizen interactions, complemented by semantic status colors (Emerald Green for completed milestones, Amber/Red for delays).", bold_prefix="Deliberate Color Hierarchy: ")
    add_bullet("Fluid layouts engineered to deliver equal clarity on 4-inch mobile smartphones in the field as on 27-inch 4K desktop monitors in municipal headquarters.", bold_prefix="Responsive Viewport Engineering: ")

    # ------------------------------------------
    # SECTION 3: CITIZEN TRANSPARENCY PORTAL
    # ------------------------------------------
    add_h1("3. Public Citizen Transparency Portal (Feature Guide)")
    add_p(
        "The public portal is accessible at https://nagaryojana.shrawankarki.com.np/ without requiring any login or credential verification. Every citizen has instantaneous access to verified municipal public works data."
    )

    add_h2("3.1 Live Metropolitan Hero Counters")
    add_p("The hero section displays four real-time municipal Key Performance Indicators (KPIs):")
    add_bullet("Aggregates municipal capital allocations across all active and completed projects, presented in authentic Nepali currency notation (e.g., रु १,२५,००,०००).", bold_prefix="Total Tracked Budget: ")
    add_bullet("Real-time count of public works currently under construction or site mobilization.", bold_prefix="Active Projects Gauge: ")
    add_bullet("Total tally of fully finished and verified civic projects delivered to the community.", bold_prefix="Completed Infrastructure: ")
    add_bullet("A dynamic city-wide completion percentage representing total capital progress across all 19 wards.", bold_prefix="Metropolitan Progress Meter: ")

    add_h2("3.2 19-Ward Interactive Municipal Jurisdiction Filter")
    add_p(
        "A horizontal scrolling filter bar features 19 individual ward chips (Ward 01 to Ward 19) along with an 'All Wards' master chip. Selecting a ward instantly isolates that jurisdiction's public works without a page reload."
    )

    add_h2("3.3 Multi-Criteria Filtering & Global Text Search")
    add_p("Citizens can combine three orthogonal search dimensions to pinpoint specific public works:")
    add_bullet("Filter across Roads & Drainage, Water & Sanitation, Community Infrastructure, Education & Sports, and Health Facilities.", bold_prefix="Sector Categories: ")
    add_bullet("Filter by Active (ongoing works), Completed (100% finished), or Delayed (behind schedule).", bold_prefix="Milestone Status: ")
    add_bullet("Live debounce-free search bar querying Project Title, Street Location, Executing Contractor / Upabhokta Samiti, and Sector simultaneously.", bold_prefix="Global Query Search: ")

    add_h2("3.4 Interactive Civic Project Cards")
    add_p("Every public project is rendered as an Apple-inspired card with comprehensive oversight data:")
    add_bullet("High-definition photographs of the project site with responsive aspect ratios and error fallbacks.", bold_prefix="Project Imagery: ")
    add_bullet("Pill badges with semantic color cues (Emerald Green for COMPLETED, Vibrant Blue for ACTIVE, Amber for DELAYED).", bold_prefix="Status Pill: ")
    add_bullet("An animated horizontal bar indicating physical progress (0% to 100%) paired with a precise percentage readout.", bold_prefix="Physical Progress Bar: ")
    add_bullet("Clear institutional tags showing Ward Number, Sector, Allocated Budget (in NPR), and Assigned Contractor / Upabhokta Samiti.", bold_prefix="Transparency Badges: ")
    add_bullet("Start dates and completion targets recorded in Bikram Sambat (B.S.) notation (e.g., '2081-11-15 (B.S.)').", bold_prefix="Bikram Sambat Calendar Dates: ")
    add_bullet("A prominent action button triggering the grievance submission modal pre-linked to that exact project.", bold_prefix="'Report Issue / गुनासो दर्ता' CTA: ")

    add_h2("3.5 Citizen Grievance Redressal Modal (ReportIssueModal.jsx)")
    add_p(
        "Public accountability requires a direct channel for residents to voice issues. Nagarयोजना implements a frictionless grievance submission modal:"
    )
    add_bullet("Zero Login Barrier: Citizens do NOT need an account, password, or verification step to submit a grievance.", bold_prefix="Universal Access: ")
    add_bullet("The modal automatically injects the Project ID, Title, and Ward into the submission payload.", bold_prefix="Context Pre-Filling: ")
    add_bullet("Citizens classify their complaint into 1 of 6 categories: Construction Delay, Substandard Material Quality, Public Safety Hazard, Transparency/Budget Discrepancy, Environmental/Noise Concerns, or Other.", bold_prefix="Structured Categorization: ")
    add_bullet("Captures the complainant's Full Name, Phone Number, Email, and Detailed Description of the problem.", bold_prefix="Contact Verification: ")
    add_bullet("Directly writes to the Firestore grievances collection with an initial PENDING status, immediately alerting the assigned Ward Executive.", bold_prefix="Instant Real-Time Persistence: ")

    add_h2("3.6 Dynamic Zero-Storage In-Browser Bilingual Localization (translator.js)")
    add_p(
        "To ensure accessibility across all demographics, the portal features seamless switching between English and authentic Nepali (नेपाली):"
    )
    add_bullet("Unlike traditional i18n implementations that require massive static JSON dictionary files (which increase bundle size and slow mobile devices), Nagarयोजना uses a zero-storage translation engine.", bold_prefix="Zero-Storage Architecture: ")
    add_bullet("Toggling the language switch sets the official googtrans=/en/ne cookie across the root path and domain, triggering the browser's native translation engine.", bold_prefix="Cookie-Driven Engine: ")
    add_bullet("The user's preference (en or ne) is stored in localStorage so their preferred language persists across sessions and page reloads.", bold_prefix="State Persistence: ")
    add_bullet("Translates Devanagari numerals, ward labels, form placeholders, and municipal civic terminology accurately.", bold_prefix="Devanagari Localization: ")

    # ------------------------------------------
    # SECTION 4: UNIFIED GATEWAY & RBAC
    # ------------------------------------------
    add_h1("4. Unified Municipal Gateway & Three-Tier Role Workspaces")
    add_p(
        "Earlier iterations of municipal software maintained fragmented, confusing login URLs (e.g., separate links for staff, contractors, and admins). Nagarयोजना consolidates all administrative access into ONE Unified Municipal Gateway (StaffPortal.jsx & StaffLoginModal.jsx)."
    )

    add_h2("4.1 Unified Login Experience & Password Toggle")
    add_bullet("Official Biratnagar City Gate insignia and municipal seal eliminate phishing risks and unify brand presentation.", bold_prefix="Branded Civic Presentation: ")
    add_bullet("Temporary generated passwords can be tricky to type on mobile touchscreens. An interactive eye toggle allows users to unmask and verify their entered password before clicking Sign In.", bold_prefix="Plaintext Password Inspection Toggle: ")
    add_bullet("Links dispatched in credential emails contain deep-linking query parameters (e.g., ?portal=staff&email=...) that automatically populate the login email input.", bold_prefix="URL Deep-Linking Auto-Fill: ")

    add_h2("4.2 Cryptographic Claims-Based Automatic Role Redirection")
    add_p(
        "Upon successful authentication via Firebase Auth, the client calls user.getIdTokenResult(true) to fetch the cryptographic custom claims embedded in the user's JWT. The application automatically redirects the authenticated user to their designated workspace:"
    )
    add_bullet("Custom claim role === 'SUPER_ADMIN' redirects to ?portal=staff (Central Metropolitan Command Center).", bold_prefix="1. Super Admin: ")
    add_bullet("Custom claim role === 'WARD_EXEC' redirects to ?portal=executive (Ward Executive Workspace).", bold_prefix="2. Ward Executive: ")
    add_bullet("Custom claim role === 'CONTRACTOR' redirects to ?portal=contractor (Civic Contractor Workspace).", bold_prefix="3. Civic Contractor: ")

    add_h2("4.3 Tier 1: Super Admin Metropolitan Command Center (?portal=staff)")
    add_p("The Super Admin workspace provides comprehensive metropolitan-wide governance across all 19 wards:")
    add_bullet("A visual interactive grid displaying every ward from Ward 01 to Ward 19. Each cell highlights whether the ward executive seat is Active (displaying appointed name, executive ID, contact, and tenure) or Vacant. Clicking an empty ward card automatically pre-selects that ward in the appointment form.", bold_prefix="19-Ward Jurisdiction Grid: ")
    add_bullet("A high-efficiency form allowing administrators to appoint Ward Executives. Requires Full Name, Official Email, Phone Number, Ward Number, and Service Period.", bold_prefix="Ward Executive Provisioning Engine: ")
    add_bullet("A custom date picker allowing administrators to select service tenure without relying on inconsistent browser-native inputs.", bold_prefix="Interactive Calendar Tenure Picker: ")
    add_bullet("Automatically computes sequential institutional identifiers like WE-4-1 (Ward 4 Executive #1) based on existing database counts.", bold_prefix="Sequential Institutional ID: ")
    add_bullet("Live roster of all active Ward Executives with real-time search, 1-click credential resend, and cascading account revocation.", bold_prefix="Executive Directory & Search: ")
    add_bullet("A dedicated dashboard tab displaying accounts locked due to brute-force attempts, showing lock expiration countdowns and providing a 1-click instant unlock button.", bold_prefix="Locked Accounts Administration: ")
    add_bullet("A metropolitan-wide user directory listing all municipal staff, ward executives, civic contractors, and administrators with role badges, ward assignments, creation dates, and last sign-in timestamps. Provides 1-click password reset dispatch via Resend and direct email outreach.", bold_prefix="Municipal Users Directory & Account Administration: ")
    add_bullet("A master switch that pauses all non-superadmin municipal operations during maintenance breaks with instant cross-tab eviction.", bold_prefix="Municipal Maintenance Break Mode: ")

    add_h2("4.4 Tier 2: Ward Executive Workspace (?portal=executive)")
    add_p("Ward Executives possess ward-isolated administration rights restricted strictly to their assigned jurisdiction:")
    add_bullet("Executives can only see, create, and manage projects and contractors belonging to their assigned wardNumber. Attempts to modify other wards are rejected by backend RBAC middleware.", bold_prefix="Strict Jurisdiction Scoping: ")
    add_bullet("Executive cards display Total Ward Budget, Active Ward Works, and Completed Infrastructure tally.", bold_prefix="Ward KPI Summary: ")
    add_bullet("Ward Executives view project progress through read-only meters, preventing administrative staff from accidentally tampering with field contractor measurements without site verification.", bold_prefix="Read-Only Progress Meters: ")
    add_bullet("A streamlined form allowing executives to launch new civic works. Captures Title, Sector, Location, Budget (in NPR), Target Completion Date (B.S.), and Project Description. Submits to POST /api/executive/create-project.", bold_prefix="'Create Infrastructure Project' Module: ")
    add_bullet("Enables executives to appoint field contractors or Upabhokta Samiti representatives for specific projects. Generates sequential IDs (e.g., CON-4-1), zero-ambiguity passwords, embeds custom claims, and triggers automated Resend credential emails.", bold_prefix="'Appoint Civic Contractor' Module: ")
    add_bullet("A dedicated inbox showing all grievances submitted by citizens for that specific ward. Executives can update complaint statuses: PENDING -> UNDER_REVIEW -> RESOLVED or REJECTED via PATCH /api/grievances/:id/status.", bold_prefix="Citizen Grievance Lifecycle Management: ")

    add_h2("4.5 Tier 3: Civic Contractor Workspace (?portal=contractor)")
    add_p(
        "Contractors, site engineers, and Upabhokta Samiti representatives operate on a focused field execution workspace:"
    )
    add_bullet("Contractors are strictly scoped to projects assigned to their UID or contractor ID, eliminating administrative clutter.", bold_prefix="Assigned Project Scoping: ")
    add_bullet("A touch-friendly range slider allowing contractors to adjust physical completion percentage from 0% to 100%.", bold_prefix="Interactive Progress Slider: ")
    add_bullet("Contractors can update the operational state to ACTIVE, COMPLETED, or DELAYED.", bold_prefix="Milestone Status Selector: ")
    add_bullet("A text area for recording physical site notes (e.g., 'Culvert slab reinforced; concrete curing in progress; awaiting gravel delivery').", bold_prefix="Field Remarks Logging: ")
    add_bullet("When a contractor saves progress, the backend updates the project record in Firestore, and the change instantly reflects on the public citizen portal and executive dashboard in real time without refreshing.", bold_prefix="Real-Time Public Broadcast: ")

    # ------------------------------------------
    # SECTION 5: ENTERPRISE SECURITY ARCHITECTURE
    # ------------------------------------------
    add_h1("5. Enterprise Security Architecture & Defense-in-Depth")
    add_p(
        "Because Nagarयोजना manages municipal public records, fiscal allocations, and citizen grievances, the architecture adheres to multi-layered defense-in-depth principles."
    )

    add_h2("5.1 3-Attempt Failed Login Lockout & IP Rate-Limiting")
    add_p("To neutralize brute-force password guessing attacks on municipal accounts:")
    add_bullet("Authentication endpoints are protected by an in-memory IP bucket rate limiter permitting a maximum of 30 requests per minute per IP address.", bold_prefix="IP Rate Limiting: ")
    add_bullet("Every failed sign-in attempt is recorded in the loginLockouts Firestore collection using an atomic database transaction. If an account logs 3 consecutive incorrect passwords, the backend immediately locks the account for 24 hours.", bold_prefix="3-Attempt Account Lockout: ")
    add_bullet("Even if an attacker discovers or guesses the valid password on the 4th attempt, the system refuses authentication and returns an explicit lockout notice.", bold_prefix="Zero-Tolerance Defense: ")
    add_bullet("Super Admins can view all currently locked emails and timestamps in their Command Center and restore access with a single click (POST /api/admin/lockouts/unlock).", bold_prefix="Super Admin 1-Click Unlock: ")
    add_bullet("Whenever an administrator or executive reprovisions an account or clicks 'Resend Credentials', the backend automatically purges prior failed lockout records.", bold_prefix="Auto-Clear on Credential Reset: ")

    add_h2("5.2 15-Minute DOM Activity Inactivity Auto-Logout")
    add_p(
        "Municipal computers and ward office terminals are often shared across multiple staff members. Leaving an active session unattended poses severe security hazards:"
    )
    add_bullet("The client continuously monitors 5 user interaction events: mousedown, keydown, scroll, touchstart, and mousemove.", bold_prefix="5-Event Activity Tracking: ")
    add_bullet("If no interaction occurs for 15 minutes (900,000 milliseconds), the client automatically invokes signOut(auth), clears sessionStorage, resets application state, and displays an informative inactivity banner.", bold_prefix="Automated Session Termination: ")

    add_h2("5.3 Session Persistence Guard (browserSessionPersistence)")
    add_p(
        "Firebase Auth is explicitly configured with browserSessionPersistence on initialization (frontend/src/firebase.js). Unlike default localStorage persistence (which keeps authentication tokens alive indefinitely even after the browser window is closed), browserSessionPersistence restricts tokens strictly to the active tab lifecycle, guaranteeing that tokens are wiped when the browser is closed."
    )

    add_h2("5.4 Municipal Maintenance Break Mode with Instant Cross-Tab Eviction")
    add_p(
        "During database migrations, major infrastructure updates, or municipal emergencies, Super Admins can toggle Municipal Maintenance Break Mode:"
    )
    add_bullet("Toggling the switch updates the systemSettings/maintenance document in Firestore.", bold_prefix="Central Master Switch: ")
    add_bullet("Every active Ward Executive and Contractor browser subscribes to this document via a real-time Firestore onSnapshot listener. The millisecond systemActive changes to false, the client immediately terminates the session (signOut(auth)), clears storage, and displays a maintenance break screen.", bold_prefix="Instant Real-Time Eviction: ")
    add_bullet("A background 3-second polling fallback and window focus event listener guarantee that even backgrounded tabs, sleeping laptops, or minimized mobile browsers are evicted the moment they are reopened.", bold_prefix="3-Second Polling Fallback: ")
    add_bullet("Non-admin users are strictly blocked at the login gateway while maintenance mode is active.", bold_prefix="Gateway Login Freeze: ")

    add_callout(
        "When a Super Admin toggles Maintenance Break, the system updates systemSettings/maintenance. All active executive and contractor browsers are subscribed to this document via Firestore onSnapshot. Upon receiving systemActive: false, the client immediately executes signOut(auth), clears sessionStorage, resets all React state hooks, and displays an informative maintenance alert. A 3-second polling fallback and window focus listener guarantee that even backgrounded tabs or sleeping devices evict immediately upon reopening.",
        title="Real-Time Maintenance Eviction Mechanism"
    )

    add_h2("5.5 Transactional Account Provisioning with Automatic Rollback")
    add_p(
        "Provisioning an institutional account involves multiple distributed network operations: creating a Firebase Auth record, writing Firestore documents, injecting custom claims, and dispatching a transactional email via Resend. If an unverified email address or network failure causes the email dispatch or database write to fail, an automatic rollback catch block immediately deletes the partial Firebase Auth record, ensuring no orphan, unusable, or zombie accounts remain in the system."
    )

    # ------------------------------------------
    # SECTION 6: AUTOMATED CREDENTIALS & EMAIL PIPELINE
    # ------------------------------------------
    add_h1("6. Automated Credentials & Email Delivery Pipeline")
    add_p(
        "Nagarयोजना implements an enterprise-grade automated email delivery pipeline using the official Resend SDK paired with DNS-verified custom domain records (admin@shrawankarki.com.np)."
    )

    add_h2("6.1 The 7-Step Transactional Dispatch Pipeline")
    add_p("When an executive or contractor is appointed, this 7-step pipeline executes behind the scenes:")
    add_bullet("Super Admin or Ward Executive inputs appointment details on the form.", bold_prefix="Step 1: Input Validation: ")
    add_bullet("Backend queries existing accounts in the ward and computes a sequential ID (e.g., WE-4-1 or CON-4-2).", bold_prefix="Step 2: Sequential ID Calculation: ")
    add_bullet("A clean 10-character password is generated using crypto.randomInt and Fisher-Yates shuffling.", bold_prefix="Step 3: Zero-Ambiguity Password: ")
    add_bullet("Firebase Auth registers the user and embeds custom claims { role, wardNumber, id }.", bold_prefix="Step 4: Firebase User Creation: ")
    add_bullet("User document is written to users/{uid} with mustChangePassword: true and status: 'active'.", bold_prefix="Step 5: Firestore Profile Storage: ")
    add_bullet("Node.js invokes the Resend SDK to dispatch an official municipal HTML card with copy-safe password boxes.", bold_prefix="Step 6: Resend API HTML Dispatch: ")
    add_bullet("If any subsequent write or email transmission fails, the partial Auth record is immediately deleted.", bold_prefix="Step 7: Automated Rollback on Error: ")

    add_h2("6.2 Zero-Ambiguity Temporary Password Generator (passwordHelper.js)")
    add_p(
        "A notorious failure mode in automated credential dispatch occurs when passwords include confusing glyphs (0, O, 1, I, l) or punctuation symbols (!@#$%&*). In mobile and webmail clients, punctuation characters are interpreted as word boundaries: when a contractor double-clicks the password in Gmail, the selection is cut off halfway, causing them to paste an incomplete password. Nagarयोजना eliminates this fundamentally through passwordHelper.js:"
    )

    add_code_block(
        "// backend/utils/passwordHelper.js\n"
        "const crypto = require('crypto');\n\n"
        "function generateTemporaryPassword(length = 10) {\n"
        "  // Excludes 0, O, 1, I, l and all punctuation symbols\n"
        "  const upper = 'ABCDEFGHJKLMNPQRSTUVWXYZ'; // 24 chars (no I, O)\n"
        "  const lower = 'abcdefghjkmnpqrstuvwxyz';  // 23 chars (no l, o)\n"
        "  const digits = '23456789';                // 8 chars (no 0, 1)\n"
        "  const all = upper + lower + digits;\n\n"
        "  const chars = [\n"
        "    upper[crypto.randomInt(0, upper.length)],\n"
        "    lower[crypto.randomInt(0, lower.length)],\n"
        "    digits[crypto.randomInt(0, digits.length)]\n"
        "  ];\n"
        "  for (let i = 3; i < length; i++) chars.push(all[crypto.randomInt(0, all.length)]);\n\n"
        "  // Fisher-Yates shuffle for uniform entropy\n"
        "  for (let i = chars.length - 1; i > 0; i--) {\n"
        "    const j = crypto.randomInt(0, i + 1);\n"
        "    [chars[i], chars[j]] = [chars[j], chars[i]];\n"
        "  }\n"
        "  return chars.join('');\n"
        "}\n\n"
        "module.exports = { generateTemporaryPassword };",
        title="Zero-Ambiguity Temporary Password Generator"
    )

    add_h2("6.3 Branded Transactional Municipal HTML Email Cards")
    add_p(
        "The Resend service dispatches bespoke HTML emails styled according to municipal design standards:"
    )
    add_bullet("Ward Executive Card: Features official Biratnagar Metropolitan insignia, personalized greeting, institutional tenure period, assigned Ward Number, Executive ID (WE-X-Y), a styled copyable password container, and a pre-linked 'Sign In to Ward Workspace' button pointing to https://nagaryojana.shrawankarki.com.np/?portal=staff.")
    add_bullet("Civic Contractor Card: Includes company/contractor name, assigned infrastructure project title, Ward Number, Contractor ID (CON-X-Y), copyable password container, and a direct 'Sign In to Field Workspace' button pointing to https://nagaryojana.shrawankarki.com.np/?portal=contractor.")
    add_bullet("Anti-Spam Domain Verification: Sent from admin@shrawankarki.com.np with validated SPF, DKIM, and DMARC DNS records, guaranteeing inbox delivery instead of spam folders.")
    add_bullet("Dispatches an Apple-styled password reset notification featuring personalized greeting, secure direct reset link button, and fall-through URL via the Resend API.", bold_prefix="Administrative Password Reset Delivery: ")
    add_bullet("Built with mobile-first CSS ensuring seamless rendering on smartphones, and click-to-copy icon badges that eliminate copy-paste friction.", bold_prefix="Mobile Responsiveness & Usability: ")

    # ------------------------------------------
    # SECTION 7: DATABASE ARCHITECTURE & SCHEMAS
    # ------------------------------------------
    add_h1("7. Database Architecture & Firestore NoSQL Schemas")
    add_p(
        "The platform utilizes Google Firebase Cloud Firestore as its primary document database. Collections are structured for rapid real-time queries, sub-second indexing, and strict role segregation."
    )

    add_h2("7.1 Firestore Collections & Document Schemas")
    add_table_data(
        headers=["Collection Path", "Key Fields & Data Types", "Description & Usage"],
        rows_data=[
            [
                "projects/{projectId}",
                "id (number|string)\ntitle (string)\nsector (string)\nlocation (string)\nbudget (string)\nprogress (number 0-100)\nstatus (ACTIVE|COMPLETED|DELAYED)\nward (string 'Ward 04')\ncontractor (string)\ncontractorUid (string)\ndate (string B.S.)\ndateLabel (string)\nimageUrl (string)\nfeatured (boolean)\nremarks (string)\nupdatedAt (timestamp)",
                "Stores all civic infrastructure projects. Read publicly by citizens; modified by Ward Executives (creation) and Civic Contractors (progress updates)."
            ],
            [
                "users/{uid}",
                "uid (string)\nemail (string)\nfullName (string)\nrole (SUPER_ADMIN|WARD_EXEC|CONTRACTOR)\nwardNumber (number 1-19)\nexecutiveId (string WE-X-Y)\ncontractorId (string CON-X-Y)\ncompanyName (string)\nprojectId (string)\nphone (string)\nservicePeriod (string)\nstatus ('active'|'revoked')\nmustChangePassword (boolean)\ncreatedAt (timestamp)",
                "Contains profile metadata and jurisdiction mappings for all authenticated personnel. Checked during login to confirm active institutional status."
            ],
            [
                "grievances/{grievanceId}",
                "id (number|string)\nprojectId (string)\nprojectTitle (string)\nward (number|string)\ncitizenName (string)\nphone (string)\nemail (string)\ncategory (string)\ndescription (string)\nstatus (PENDING|UNDER_REVIEW|RESOLVED|REJECTED)\ncreatedAt (timestamp)",
                "Stores public grievances submitted through the Report Issue modal. Filtered by ward for executive oversight and resolution."
            ],
            [
                "systemSettings/maintenance",
                "systemActive (boolean)\nmaintenanceMode (boolean)\nupdatedBy (string)\nupdatedAt (timestamp)",
                "Single document controlling metropolitan system availability. Listened to by all active clients for instant cross-tab session eviction."
            ],
            [
                "loginLockouts/{email}",
                "failedAttempts (number 0-3)\nlockedUntil (timestamp)\nlastAttemptAt (timestamp)",
                "Tracks consecutive failed login attempts per email address. Enforces the 24-hour lockout window upon reaching 3 failed attempts."
            ]
        ],
        col_widths=[Inches(1.8), Inches(2.6), Inches(2.6)]
    )

    add_h2("7.2 Firestore Security Rules (firestore.rules)")
    add_p(
        "Firestore security rules enforce authorization directly at the database layer:"
    )
    add_bullet("Public Read Access: The projects and systemSettings collections permit universal read access (allow read: if true;) to power the citizen portal and maintenance listeners.")
    add_bullet("Citizen Grievance Submission: The grievances collection allows unauthenticated write access (allow create: if true;) so any resident can file a complaint without account friction.")
    add_bullet("Restricted Modifications: Modifying projects, resolving grievances, or reading users requires valid Firebase ID tokens containing authorized custom claims.")

    # ------------------------------------------
    # SECTION 8: RESTFUL API ENDPOINT SPECIFICATION
    # ------------------------------------------
    add_h1("8. RESTful API Endpoint Reference Guide")
    add_p(
        "The Express.js REST API runs at https://nagar-yojana.onrender.com. All administrative routes require an Authorization: Bearer <ID_TOKEN> header validated by verifyToken middleware."
    )

    add_table_data(
        headers=["HTTP Method & Route", "Authorization Level", "Request / Response Summary", "Purpose & Operational Details"],
        rows_data=[
            [
                "GET /",
                "Public (None)",
                "Res: 200 OK ('NagarYojana API is awake and running!')",
                "Root health check endpoint pinged by UptimeRobot every 5 minutes to prevent cold spin-down."
            ],
            [
                "GET /health",
                "Public (None)",
                "Res: 200 OK ({ status: 'healthy', timestamp })",
                "Secondary health check endpoint providing server uptime verification."
            ],
            [
                "GET /api/projects",
                "Public (None)",
                "Query: ?ward=&status=&sector=\nRes: 200 OK ([ Project ])",
                "Fetches project dataset for public portal, filtered by ward, milestone status, or sector."
            ],
            [
                "POST /api/grievances",
                "Public (Citizen)",
                "Body: { projectId, citizenName, phone, email, category, description }\nRes: 201 Created",
                "Public endpoint for citizens to submit grievances directly from the Report Issue modal."
            ],
            [
                "GET /api/auth/system-status",
                "Public (None)",
                "Res: 200 OK ({ systemActive: boolean })",
                "Returns maintenance mode status for login gateway checks."
            ],
            [
                "POST /api/auth/lockout-status",
                "Public (Rate Limited)",
                "Body: { email }\nRes: 200 OK ({ locked: boolean, attempts: number })",
                "Checks if a specific email address is currently locked out before attempting authentication."
            ],
            [
                "POST /api/auth/record-failure",
                "Public (Rate Limited)",
                "Body: { email }\nRes: 200 OK ({ locked: boolean, failedAttempts: number })",
                "Increments failed attempt count in loginLockouts; locks account for 24h upon 3rd failure."
            ],
            [
                "POST /api/auth/clear-failures",
                "Authenticated User",
                "Res: 200 OK ({ cleared: true })",
                "Clears failed attempt counters upon successful authenticated login."
            ],
            [
                "POST /api/admin/provision-ward-exec",
                "SUPER_ADMIN",
                "Body: { fullName, email, phone, wardNumber, servicePeriod }\nRes: 201 Created",
                "Generates WE-X-Y ID, clean password, creates Auth + Firestore, and sends Resend credentials email."
            ],
            [
                "GET /api/admin/ward-execs",
                "WARD_EXEC | SUPER_ADMIN",
                "Res: 200 OK ([ ExecutiveProfile ])",
                "Retrieves directory of all appointed Ward Executives across all 19 wards."
            ],
            [
                "GET /api/admin/next-id/:wardNumber",
                "WARD_EXEC | SUPER_ADMIN",
                "Res: 200 OK ({ nextExecutiveId: string })",
                "Pre-calculates the next sequential ID for a given ward (e.g. WE-4-2)."
            ],
            [
                "GET /api/admin/users",
                "SUPER_ADMIN",
                "Res: 200 OK ({ users: [ MunicipalUser ] })",
                "Super Admin directory aggregating Firestore user profiles and Firebase Auth metadata with last active sign-in timestamps."
            ],
            [
                "POST /api/admin/users/reset-password",
                "SUPER_ADMIN",
                "Body: { email, name }\nRes: 200 OK ({ success, resetLink })",
                "Generates cryptographic password reset link and dispatches branded email via Resend API."
            ],
            [
                "GET /api/admin/lockouts",
                "SUPER_ADMIN",
                "Res: 200 OK ([ LockoutRecord ])",
                "Lists all currently locked email addresses with timestamps for Super Admin review."
            ],
            [
                "POST /api/admin/lockouts/unlock",
                "SUPER_ADMIN",
                "Body: { email }\nRes: 200 OK ({ unlocked: true })",
                "Super Admin 1-click unlock: removes the lockout record from loginLockouts immediately."
            ],
            [
                "DELETE /api/admin/ward-exec/:uid",
                "SUPER_ADMIN",
                "Query: ?email=\nRes: 200 OK ({ deleted: true })",
                "Cascading revocation: deletes Auth account, user Firestore doc, and cleans stray records."
            ],
            [
                "POST /api/admin/ward-exec/:uid/resend",
                "SUPER_ADMIN",
                "Res: 200 OK ({ resent: true })",
                "Regenerates a clean temporary password, updates Auth, and re-dispatches the email card."
            ],
            [
                "POST /api/admin/toggle-system-status",
                "SUPER_ADMIN",
                "Body: { systemActive: boolean }\nRes: 200 OK ({ systemActive })",
                "Toggles Municipal Maintenance Break Mode, triggering instant cross-tab eviction."
            ],
            [
                "POST /api/executive/create-project",
                "WARD_EXEC",
                "Body: { title, sector, location, budget, targetDate, description }\nRes: 201 Created",
                "Ward Executive project inception module. Creates project locked to the executive's ward."
            ],
            [
                "POST /api/executive/provision-contractor",
                "WARD_EXEC",
                "Body: { name, email, phone, projectId }\nRes: 201 Created",
                "Appoints civic contractor, generates CON-X-Y ID and clean password, sends Resend email."
            ],
            [
                "GET /api/executive/contractors",
                "WARD_EXEC",
                "Res: 200 OK ([ ContractorProfile ])",
                "Retrieves list of contractors appointed within the executive's ward jurisdiction."
            ],
            [
                "POST /api/executive/contractor/:uid/resend",
                "WARD_EXEC",
                "Res: 200 OK ({ resent: true })",
                "Regenerates credentials and re-dispatches notification email to a civic contractor."
            ],
            [
                "DELETE /api/executive/contractor/:uid",
                "WARD_EXEC",
                "Res: 200 OK ({ revoked: true })",
                "Revokes contractor access, deleting Auth credentials and marking profile revoked."
            ],
            [
                "GET /api/grievances",
                "WARD_EXEC",
                "Query: ?ward=\nRes: 200 OK ([ Grievance ])",
                "Fetches citizen grievances submitted for the executive's specific ward jurisdiction."
            ],
            [
                "PATCH /api/grievances/:id/status",
                "WARD_EXEC",
                "Body: { status: 'UNDER_REVIEW'|'RESOLVED'|'REJECTED' }\nRes: 200 OK",
                "Updates the lifecycle status of a citizen complaint."
            ],
            [
                "PUT /api/projects/:id",
                "CONTRACTOR | WARD_EXEC",
                "Body: { progress: number, status: string, remarks: string }\nRes: 200 OK",
                "Field contractor milestone update endpoint. Saves progress, status, and engineering remarks."
            ]
        ],
        col_widths=[Inches(1.8), Inches(1.3), Inches(2.1), Inches(1.8)]
    )

    # ------------------------------------------
    # SECTION 9: DEPLOYMENT & PRODUCTION RUNBOOK
    # ------------------------------------------
    add_h1("9. Production Deployment & Operations Runbook")
    add_p(
        "Nagarयोजना operates on a high-availability dual-cloud architecture. Static client assets are served from global edge caches, while API execution is handled in containerized cloud runtimes."
    )

    add_h2("9.1 Dual-Cloud Hosting Architecture")
    add_bullet("Frontend Edge Hosting (Cloudflare Pages): Hosts the compiled production React bundle (dist/). Assets are served across hundreds of worldwide edge data centers with instant caching, Brotli compression, and sub-100ms load times worldwide.")
    add_bullet("Backend API Hosting (Render Web Service): Runs the Node.js and Express backend at https://nagar-yojana.onrender.com. Connects securely to the Firebase Admin SDK using base64-encoded service account credentials stored in environment variables.")
    add_bullet("Dynamic Environment Detection (apiConfig.js): Automatically routes API calls to https://nagar-yojana.onrender.com in production and localhost:5000 during local development.")
    add_bullet("Cold-Start Neutralization via UptimeRobot: Cloud container free tiers spin down after 15 minutes of idle time. External monitoring via UptimeRobot pings GET / and GET /health every 5 minutes, keeping the container warm 24/7.")

    add_h2("9.2 Step-by-Step Replication Roadmap (From Scratch to Production)")
    add_p("To replicate this complete civic infrastructure system from scratch, follow this 6-step roadmap:")
    add_bullet("Create root directory with frontend/ and backend/ folders. Configure root package.json with concurrently to start Vite (port 5173) and Express (port 5000) simultaneously via 'npm run dev'.", bold_prefix="Step 1: Monorepo Setup & Concurrency: ")
    add_bullet("Create a project at console.firebase.google.com. Enable Firebase Authentication (Email/Password) and Cloud Firestore. Generate serviceAccountKey.json for backend/firebaseAdmin.js and client credentials for frontend/src/firebase.js.", bold_prefix="Step 2: Firebase Admin & Client Initialization: ")
    add_bullet("Implement modular Express routes: /api/admin, /api/executive, /api/auth, /api/projects, and /api/grievances. Secure administrative routes with verifyToken and role middleware.", bold_prefix="Step 3: Build Express REST Endpoints with RBAC: ")
    add_bullet("Verify custom domain DNS records (SPF, DKIM, DMARC) in Resend. Install the resend SDK in backend. Build responsive HTML email templates with copyable password boxes.", bold_prefix="Step 4: Configure Resend Transactional Email: ")
    add_bullet("Build responsive UI components using Tailwind CSS: Hero banner, 19-Ward Filter chips, Report Issue modal, Unified Login, and 3-Tier role dashboards. Use onSnapshot for live Firestore sync.", bold_prefix="Step 5: Design Apple-Inspired React Interface: ")
    add_bullet("Deploy React to Cloudflare Pages (Build command: npm run build, Directory: dist). Deploy Express to Render. Configure UptimeRobot to ping GET / every 5 minutes.", bold_prefix="Step 6: Deploy to Dual Cloud & Configure Uptime: ")

    add_h2("9.3 Maintenance & Emergency Recovery Runbook")
    add_bullet("If a municipal officer is locked out after 3 incorrect attempts, navigate to ?portal=staff -> Locked Accounts tab -> click 'Unlock Account'. Access is restored immediately without waiting 24 hours.", bold_prefix="Unlocking a Locked Account: ")
    add_bullet("If an executive or contractor misplaces their password, navigate to the Executive or Contractor Directory and click 'Resend Credentials'. A fresh zero-ambiguity password is generated and emailed immediately.", bold_prefix="Credential Recovery: ")
    add_bullet("To perform database maintenance, flip the 'Maintenance Break' toggle in the Super Admin dashboard. All executive and contractor sessions across all open tabs are evicted within milliseconds.", bold_prefix="Freezing Operations During Maintenance: ")

    # ------------------------------------------
    # SECTION 10: VISUAL ARCHITECTURE & INTERFACE GALLERY
    # ------------------------------------------
    add_h1("10. Visual Architecture & Interface Gallery")
    add_p("The following high-resolution figures illustrate the core interfaces and operational workflows of the live production system:")

    add_image_figure(
        "image1.png",
        "Figure 1: Super Admin Command Center & Provisioning Workspace (Ward Executive Appointments & Real-Time Directory)"
    )
    add_p(
        "Figure 1 illustrates the Super Admin Municipal Command Center (?portal=staff). The interface features the centralized executive provisioning workspace: administrators can select a target ward (1–19), input the official's name, verified email, phone number, and service tenure (via the built-in calendar picker). Below the form sits the real-time Executive Directory displaying institutional IDs (WE-X-Y), ward numbers, appointment dates, and 1-click action triggers to Resend Credentials or Revoke access."
    )

    add_image_figure(
        "image2.png",
        "Figure 2: 19-Ward Municipal Jurisdiction Grid (Visual Status of Active and Vacant Ward Executive Seats)"
    )
    add_p(
        "Figure 2 displays the interactive 19-Ward Municipal Jurisdiction Grid. Every ward from Ward 01 to Ward 19 is represented by a dedicated status card. Green indicators highlight actively occupied executive seats (displaying the assigned executive's name, tenure, and contact), while vacant ward seats show an immediate 'Appoint Executive' call-to-action button that pre-populates the appointment form."
    )

    add_image_figure(
        "image3.png",
        "Figure 3: Civic Infrastructure Public Explorer & Citizen Transparency Portal (Biratnagar Metropolitan City)"
    )
    add_p(
        "Figure 3 demonstrates the citizen-facing public portal (https://nagaryojana.shrawankarki.com.np/). The view showcases the dynamic 19-ward filter bar, real-time capital budget counters, sector filter chips, and Apple-inspired civic project cards. Each card displays verified physical progress bars (0–100%), milestone status badges (COMPLETED, ACTIVE, DELAYED), contractor attributions, allocated budgets in Nepali currency notation, and the direct 'Report Issue / गुनासो दर्ता' citizen action button."
    )

    # ------------------------------------------
    # SECTION 11: FUTURE ROADMAP & MUNICIPAL EXPANSION
    # ------------------------------------------
    add_h1("11. Future Roadmap & Municipal Expansion")
    add_p(
        "As Nagarयोजना expands across additional municipalities throughout Nepal, the engineering roadmap includes several planned capabilities:"
    )
    add_bullet("Integration of Leaflet/Mapbox vector polygons showing exact ward boundaries and geotagged project pinpoints across Biratnagar.", bold_prefix="1. Geospatial GIS Interactive Ward Mapping: ")
    add_bullet("Enhancing the citizen Report Issue modal to accept smartphone photos and videos of construction defects, uploaded securely to Firebase Storage with automated image compression.", bold_prefix="2. Multimedia Citizen Grievance Uploads: ")
    add_bullet("Enabling local consumer committees (*Upabhokta Samiti*) to upload verified vendor invoices and expense receipts, generating automated public audit ledgers.", bold_prefix="3. Upabhokta Samiti Public Expense Ledgers: ")
    add_bullet("Direct SMS notifications to ward residents and contractors via local telecom gateways (Sparrow SMS / Nepal Telecom) when milestone status changes occur.", bold_prefix="4. Real-Time SMS Municipal Bulletins: ")

    # Save to target destination
    output_path = os.path.join("docs", "overalldocx.docx")
    doc.save(output_path)
    print(f"Master document successfully written to: {output_path}")

if __name__ == "__main__":
    create_document()
